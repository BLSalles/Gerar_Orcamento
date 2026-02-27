from flask import Flask, request, render_template, send_from_directory, url_for
from docx import Document
from datetime import datetime
import os
import re

app = Flask(__name__)

# Garante que a pasta de saída exista (importante no Render)
OUTPUT_DIR = os.path.join("static", "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)


@app.route("/")
def index():
    return render_template("index.html")


def formatar_valor(valor: float) -> str:
    return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def nome_arquivo_seguro(nome: str) -> str:
    """
    Remove caracteres perigosos para nome de arquivo e evita path traversal.
    """
    if not nome:
        return "CLIENTE"
    nome = nome.strip()

    # Remove caracteres inválidos em nomes de arquivo (Windows/Linux)
    nome = re.sub(r'[\\/:*?"<>|]+', "", nome)

    # Opcional: troca múltiplos espaços por 1
    nome = re.sub(r"\s+", " ", nome)

    # Se sobrar vazio, usa fallback
    return nome or "CLIENTE"


def parse_data_form(campo: str) -> str:
    """
    Recebe data do input type=date (YYYY-MM-DD) e retorna DD/MM/YYYY.
    Se vier vazio, retorna 'Data não informada'.
    """
    if not campo:
        return "Data não informada"
    try:
        return datetime.strptime(campo, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return "Data inválida"


@app.route("/download/<path:filename>")
def download_file(filename):
    # Segurança extra: só permite baixar arquivo dentro da pasta OUTPUT_DIR
    # send_from_directory já ajuda, mas vamos garantir que não tenha barras.
    filename = os.path.basename(filename)
    return send_from_directory(OUTPUT_DIR, filename, as_attachment=True)


@app.route("/atualizar_word", methods=["POST"])
def atualizar_word():
    try:
        nome_cliente = request.form.get("nome", "Cliente não informado")
        endereco = request.form.get("endereco", "Endereço não informado")
        dias_raw = request.form.get("dias", "0")
        numero_pedido = request.form.get("n_pedido", "Pedido não informado")
        forma_pagamento = request.form.get("forma_pagamento", "Forma de pagamento não informada")

        # Datas
        data_atual = datetime.now().strftime("%d/%m/%Y")
        data_prevista = parse_data_form(request.form.get("data_prevista"))
        data_vencimento = parse_data_form(request.form.get("data_vencimento"))

        # Dias
        try:
            dias = int(dias_raw)
        except ValueError:
            return "Erro: 'Dias' deve ser um número inteiro válido.", 400

        # Itens
        codigos = request.form.getlist("codigo[]")
        descricoes = request.form.getlist("descricao[]")
        unidades = request.form.getlist("un[]")
        quantidades = request.form.getlist("qtd[]")
        valores_unitarios = request.form.getlist("valor_unit[]")

        if not (len(codigos) == len(descricoes) == len(unidades) == len(quantidades) == len(valores_unitarios)):
            return "Erro: Todos os itens devem ser preenchidos e ter o mesmo comprimento.", 400

        itens = []
        numero_itens = 0
        soma_quantidades = 0
        total_produtos = 0.0

        for i, (codigo, descricao, un, qtd, valor_unit) in enumerate(
            zip(codigos, descricoes, unidades, quantidades, valores_unitarios), start=1
        ):
            try:
                qtd_int = int(qtd)
                valor_float = float(valor_unit.replace(",", "."))  # aceita "10,50" também
                total = qtd_int * valor_float

                itens.append(
                    (
                        codigo,
                        descricao,
                        un,
                        qtd_int,
                        formatar_valor(valor_float),
                        formatar_valor(total),
                    )
                )

                numero_itens += 1
                soma_quantidades += qtd_int
                total_produtos += total
            except ValueError:
                return (
                    f"Erro: Dados inválidos no item {i}. "
                    "Certifique-se de que as quantidades e valores unitários sejam numéricos.",
                    400,
                )

        total_pedido = total_produtos

        # Arquivos
        arquivo_entrada = "ORÇAMENTO_LUCAS.docx"

        # Nome seguro para arquivo final
        cliente_safe = nome_arquivo_seguro(nome_cliente)
        nome_saida = f"ORÇAMENTO_{cliente_safe}.docx"
        arquivo_saida = os.path.join(OUTPUT_DIR, nome_saida)

        # Edita e salva
        editar_word(
            arquivo_entrada=arquivo_entrada,
            arquivo_saida=arquivo_saida,
            nome=nome_cliente,
            endereco=endereco,
            dias=dias,
            numero_pedido=numero_pedido,
            data_atual=data_atual,
            data_prevista=data_prevista,
            data_vencimento=data_vencimento,
            forma_pagamento=forma_pagamento,
            itens=itens,
            numero_itens=numero_itens,
            soma_quantidades=soma_quantidades,
            total_pedido=formatar_valor(total_pedido),
            total_produtos=formatar_valor(total_produtos),
        )

        # ✅ LINK CORRETO: passa somente o nome do arquivo (não o caminho todo)
        download_url = url_for("download_file", filename=nome_saida)

        return f"Documento Word atualizado com sucesso! Faça o download <a href='{download_url}'>aqui</a>."

    except Exception as e:
        return f"Erro ao processar o documento Word: {e}", 500


def substituir_em_paragrafos(doc: Document, mapa: dict):
    """
    Substitui placeholders em parágrafos (simples).
    OBS: se seus placeholders estiverem quebrados em "runs", isso pode não pegar.
    """
    for p in doc.paragraphs:
        for chave, valor in mapa.items():
            if chave in p.text:
                p.text = p.text.replace(chave, str(valor))


def editar_word(
    arquivo_entrada,
    arquivo_saida,
    nome,
    endereco,
    dias,
    numero_pedido,
    data_atual,
    data_prevista,
    data_vencimento,
    forma_pagamento,
    itens,
    numero_itens,
    soma_quantidades,
    total_pedido,
    total_produtos,
):
    try:
        doc = Document(arquivo_entrada)

        # Exemplo de placeholders (ajuste conforme seu DOCX)
        mapa = {
            "[nome]": nome,
            "[endereco]": endereco,
            "[dias]": dias,
            "[n_pedido]": numero_pedido,
            "[data_atual]": data_atual,
            "[data_prevista]": data_prevista,
            "[data_vencimento]": data_vencimento,
            "[forma_pagamento]": forma_pagamento,
            "[numero_itens]": numero_itens,
            "[soma_quantidades]": soma_quantidades,
            "[total_pedido]": total_pedido,
            "[total_produtos]": total_produtos,
        }

        substituir_em_paragrafos(doc, mapa)

        # Se você tem tabela de itens no Word, aqui é onde você deve preencher.
        # (Mantive sua lógica como placeholder, porque depende do seu template.)

        doc.save(arquivo_saida)

    except Exception as e:
        raise RuntimeError(f"Erro no processo de edição do Word: {e}")


if __name__ == "__main__":
    app.run(debug=True)